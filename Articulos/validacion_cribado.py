"""Validación del cribado por dos revisores independientes (sustituye a Rayyan).

Paso 1 — generar las planillas ciegas (una por revisor):
    python Articulos/validacion_cribado.py generar
  Crea Articulos/validacion_cribado_andres.xlsx y Articulos/validacion_cribado_cristian.xlsx con
  todos los estudios incluidos y una muestra aleatoria del 20 % de los excluidos (semilla fija),
  SIN la decisión del primer revisor. Cada planilla tiene un orden distinto, para que cada revisor
  trabaje por su cuenta. Cada uno llena "Decisión" y "Código" con las listas desplegables.

Paso 2 — calcular el acuerdo:
    python Articulos/validacion_cribado.py kappa
  Lee las planillas llenas, compara cada revisor con el primer cribado (Articulos/cribado.csv) y
  entre sí (kappa de Cohen), calcula el kappa de Fleiss de los tres, y guarda en
  Articulos/desacuerdos_cribado.csv los registros con desacuerdo y la decisión por mayoría.
"""
import csv
import os
import random
import re
import sys
from collections import Counter

sys.stdout.reconfigure(encoding="utf-8")
BASE = os.path.dirname(os.path.abspath(__file__))
REVISORES = {"Andrés": "validacion_cribado_andres.xlsx", "Cristian": "validacion_cribado_cristian.xlsx"}
CRITERIOS = [
    ("CI1", "Estudio primario que clasifica texto de requerimientos en tipos (FR/NFR, subclases) o atributos de calidad"),
    ("CI2", "Usa ML, aprendizaje profundo, transformers o LLM"),
    ("CI3", "Reporta evaluación empírica cuantitativa (exactitud, precisión, recall, F1…)"),
    ("CE1", "Fuera del dominio de software (manufactura, medicina, redes…)"),
    ("CE2", "Otra tarea de software/RE (ambigüedad, trazabilidad, generación, bugs…) o clases que no son tipos de requerimiento"),
    ("CE3", "Sin evaluación empírica (visión, tutorial, dataset sin experimentos)"),
    ("CE4", "Estudio secundario (revisión, mapeo, survey)"),
    ("CE5", "Duplicado o versión resumida/extendida de otro registro"),
]


def leer_cribado():
    with open(os.path.join(BASE, "cribado.csv"), encoding="utf-8-sig") as fh:
        return {int(r["id"]): r for r in csv.DictReader(fh)}


def resumenes():
    out = {}
    for b in ["S1", "S2", "S3", "S4"]:
        with open(os.path.join(BASE, "SCUPUS", f"{b}.csv"), encoding="utf-8-sig") as fh:
            for x in csv.DictReader(fh):
                out.setdefault((x["DOI"] or x["Title"]).strip().lower(), re.sub(r"\s+", " ", x["Abstract"]))
    return out


def muestra():
    crib = leer_cribado()
    incluidos = [i for i, r in crib.items() if r["decision"] == "I"]
    excluidos = [i for i, r in crib.items() if r["decision"] != "I"]
    return incluidos + random.Random(2026).sample(excluidos, round(0.2 * len(excluidos)))


def generar():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    crib, abst, filas = leer_cribado(), resumenes(), muestra()
    for n, (nombre, archivo) in enumerate(REVISORES.items()):
        orden = list(filas)
        random.Random(100 + n).shuffle(orden)
        wb = Workbook()
        ws = wb.active
        ws.title = "Cribado"
        ws.append(["ID", "Año", "Título", "Resumen", "Decisión (I/E)", "Código (CE1–CE5 si E)", "Comentario"])
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F4E78")
        for i in orden:
            r = crib[i]
            ws.append([i, int(r["anio"]), r["titulo"], abst.get((r["doi"] or r["titulo"]).strip().lower(), ""),
                       None, None, None])
        for col, ancho in zip("ABCDEFG", [6, 7, 45, 90, 14, 20, 30]):
            ws.column_dimensions[col].width = ancho
        for fila in ws.iter_rows(min_row=2):
            for c in fila:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        fin = len(orden) + 1
        dv1 = DataValidation(type="list", formula1='"I,E"', allow_blank=True)
        dv2 = DataValidation(type="list", formula1='"CE1,CE2,CE3,CE4,CE5"', allow_blank=True)
        ws.add_data_validation(dv1)
        ws.add_data_validation(dv2)
        dv1.add(f"E2:E{fin}")
        dv2.add(f"F2:F{fin}")
        for fila in ws[f"E2:F{fin}"]:
            for x in fila:
                x.fill = PatternFill("solid", fgColor="FFF2CC")
        ws.freeze_panes = "C2"

        inst = wb.create_sheet("Instrucciones")
        inst.append([f"Planilla de {nombre} — revisión ciega del cribado"])
        inst["A1"].font = Font(bold=True, size=13)
        for t in ["1. Trabaja solo: no compartas tus respuestas ni mires las del otro revisor.",
                  "2. No abras cribado.csv ni incluidos.csv: tienen las decisiones del primer revisor.",
                  "3. Lee el título y el resumen de cada fila.",
                  "4. En 'Decisión' elige I (incluir: cumple CI1, CI2 y CI3) o E (excluir).",
                  "5. Si eliges E, indica en 'Código' el criterio de exclusión que aplica.",
                  "6. Regla de duda: si el resumen no basta para decidir, marca I (pasa a texto completo).",
                  f"7. Guarda el archivo como .xlsx con el mismo nombre ({archivo}) y envíaselo a Alessandro.", ""]:
            inst.append([t])
        inst.append(["Código", "Criterio"])
        for c in inst[inst.max_row]:
            c.font = Font(bold=True)
        for cod, txt in CRITERIOS:
            inst.append([cod, txt])
        inst.column_dimensions["A"].width = 12
        inst.column_dimensions["B"].width = 110
        wb.move_sheet("Instrucciones", offset=-1)
        wb.save(os.path.join(BASE, archivo))
        print(f"Planilla de {nombre}: Articulos/{archivo} ({len(orden)} filas)")


def leer_planilla(archivo):
    from openpyxl import load_workbook
    ws = load_workbook(os.path.join(BASE, archivo))["Cribado"]
    out = {}
    for i, _, titulo, _, dec, cod, com in ws.iter_rows(min_row=2, values_only=True):
        if dec in ("I", "E"):
            out[i] = (dec, cod or "", com or "", titulo)
    return out


def cohen(a, b):
    n = len(a)
    po = sum(x == y for x, y in zip(a, b)) / n
    ca, cb = Counter(a), Counter(b)
    pe = sum(ca[k] * cb[k] for k in ("I", "E")) / n ** 2
    return po, ((po - pe) / (1 - pe) if pe < 1 else 1.0)


def fleiss(tabla):
    """tabla: lista de listas con las decisiones de cada evaluador por registro."""
    n, m = len(tabla), len(tabla[0])
    p_cat = {k: sum(f.count(k) for f in tabla) / (n * m) for k in ("I", "E")}
    p_i = [(sum(f.count(k) ** 2 for k in ("I", "E")) - m) / (m * (m - 1)) for f in tabla]
    pbar, pe = sum(p_i) / n, sum(v ** 2 for v in p_cat.values())
    return (pbar - pe) / (1 - pe) if pe < 1 else 1.0


def nivel(k):
    return ("casi perfecto" if k > .8 else "sustancial" if k > .6 else "moderado" if k > .4
            else "aceptable" if k > .2 else "leve")


def kappa():
    crib = leer_cribado()
    r1 = {i: ("I" if crib[i]["decision"] == "I" else "E") for i in muestra()}
    revs = {}
    for nombre, archivo in REVISORES.items():
        if os.path.exists(os.path.join(BASE, archivo)):
            d = leer_planilla(archivo)
            print(f"{nombre}: {len(d)} de {len(r1)} filas con decisión")
            if d:
                revs[nombre] = d
    if not revs:
        sys.exit("Todavía no hay planillas llenas.")

    print("\nKappa de Cohen por pares")
    evaluadores = {"Primer revisor": r1, **{k: {i: v[0] for i, v in d.items()} for k, d in revs.items()}}
    nombres = list(evaluadores)
    for x in range(len(nombres)):
        for y in range(x + 1, len(nombres)):
            a, b = evaluadores[nombres[x]], evaluadores[nombres[y]]
            comunes = sorted(set(a) & set(b))
            if not comunes:
                continue
            po, k = cohen([a[i] for i in comunes], [b[i] for i in comunes])
            print(f"  {nombres[x]} vs {nombres[y]}: n={len(comunes)}  acuerdo={po:.1%}  "
                  f"kappa={k:.3f} ({nivel(k)})")

    todos = sorted(set.intersection(*[set(e) for e in evaluadores.values()]))
    if len(evaluadores) >= 3 and todos:
        kf = fleiss([[evaluadores[n][i] for n in nombres] for i in todos])
        print(f"\nKappa de Fleiss ({len(nombres)} evaluadores, n={len(todos)}): {kf:.3f} ({nivel(kf)})")

    filas, cambios = [], 0
    for i in todos:
        votos = [evaluadores[n][i] for n in nombres]
        if len(set(votos)) > 1:
            mayoria = Counter(votos).most_common(1)[0][0]
            cambios += mayoria != r1[i]
            filas.append([i, crib[i]["titulo"], *votos, mayoria,
                          " | ".join(f"{n}: {revs[n][i][1]} {revs[n][i][2]}".strip() for n in revs)])
    with open(os.path.join(BASE, "desacuerdos_cribado.csv"), "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        w.writerow(["id", "titulo", *nombres, "mayoria", "codigos_y_comentarios"])
        w.writerows(filas)
    print(f"\n{len(filas)} registros con algún desacuerdo -> Articulos/desacuerdos_cribado.csv")
    print(f"En {cambios} de ellos la mayoría contradice al primer revisor (revisarlos juntos).")


if __name__ == "__main__":
    {"generar": generar, "kappa": kappa}.get(sys.argv[1] if len(sys.argv) > 1 else "", lambda: print(__doc__))()
