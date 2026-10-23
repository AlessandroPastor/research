"""Crea RSL/formulario_extraccion_calidad.xlsx: formulario de extracción de datos y evaluación de calidad.

Hojas:
  Instrucciones  cómo se llena cada hoja
  Extraccion     una fila por estudio incluido; columnas grises = sugerencia automática (verificar)
  Calidad        lista de 8 preguntas (QA1–QA8) por estudio, con puntaje total automático
  Criterios      definición de cada pregunta de calidad y de cada campo de extracción

La lista de calidad adapta los criterios de Dybå y Dingsøyr (2008) y las guías de Kitchenham y
Charters (2007) al problema de la clasificación de requerimientos con desbalance de clases.

Uso:  python Articulos/formulario_rsl.py   (no sobrescribe un formulario ya existente)
"""
import csv
import os
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

BASE = os.path.dirname(os.path.abspath(__file__))
SALIDA = os.path.join(BASE, "..", "RSL", "formulario_extraccion_calidad.xlsx")
if os.path.exists(SALIDA) and "--forzar" not in sys.argv:
    sys.exit("El formulario ya existe; no se sobrescribe (use --forzar si de verdad quiere regenerarlo).")

AZUL, GRIS, AMARILLO = "1F4E78", "EDEDED", "FFF2CC"

CAMPOS = [  # (columna, descripción, valores permitidos o None)
    ("Fuente del texto", "SRS, historias de usuario, reseñas, issues, otro", "SRS,Historias de usuario,Reseñas/feedback,Issues/GitHub,Otro"),
    ("Conjunto(s) de datos", "Nombre y versión: PROMISE, PROMISE_exp, PURE, SecReq, propio…", None),
    ("Idioma", "Idioma de los requerimientos", "Inglés,Español,Otro,Varios"),
    ("N.º de requerimientos", "Tamaño total del conjunto usado", None),
    ("Tarea", "Tipo de clasificación", "Binaria FR/NFR,Multiclase NFR,Multiclase FR+NFR,Multietiqueta,Jerárquica,Otra"),
    ("N.º de clases", "Cuántas clases se clasifican", None),
    ("¿Reporta la distribución de clases?", "Si informa cuántos ejemplos hay por clase", "Sí,No"),
    ("Técnica contra el desbalance", "SMOTE, SMOTE-Tomek, EDA, TextAttack, retrotraducción, LLM, pérdida, ninguna…", None),
    ("¿Aumenta solo el entrenamiento?", "Si la aumentación se aplica después de separar la prueba", "Sí,No,No se indica,No aplica"),
    ("Clasificador(es)", "SVM, CNN, BiLSTM, BERT, RoBERTa, GPT-4o, Llama…", None),
    ("Uso del LLM", "Papel del LLM en el estudio", "Clasificador zero-shot,Clasificador few-shot,Ajuste fino,Generador de datos,No usa LLM"),
    ("Protocolo de evaluación", "Partición fija (indicar %), validación cruzada (k), entre conjuntos…", None),
    ("¿Controla duplicados entre entrenamiento y prueba?", "Si menciona deduplicación o fuga de datos", "Sí,No,No se indica"),
    ("Métricas reportadas", "Exactitud, precisión, exhaustividad, F1 macro/ponderado/por clase, MCC…", None),
    ("F1 macro", "Mejor valor reportado (0–1)", None),
    ("F1 ponderado", "Mejor valor reportado (0–1)", None),
    ("F1 de clases minoritarias", "Valores para PO, FT, L, MN u otras clases raras, si se reportan", None),
    ("Línea base comparada", "Contra qué se compara: sin aumentación, NoRBERT, SVM, otro LLM…", None),
    ("¿Prueba estadística?", "Si usa pruebas de significancia o reporta variabilidad", "Sí,No"),
    ("¿Código o datos disponibles?", "Repositorio público o anexo", "Sí,No,Parcial"),
    ("Limitaciones declaradas", "Resumen breve", None),
    ("Aporte a la PR", "PR1–PR5 a las que aporta evidencia", None),
    ("Notas del revisor", "Observaciones", None),
]
CALIDAD = [
    ("QA1", "¿Los objetivos o preguntas de investigación están claramente definidos?"),
    ("QA2", "¿El conjunto de datos está descrito (origen, tamaño, clases y su distribución)?"),
    ("QA3", "¿El protocolo de evaluación está descrito y evita la fuga de datos (la aumentación solo en "
            "entrenamiento, sin duplicados entre entrenamiento y prueba)?"),
    ("QA4", "¿Usa métricas adecuadas para el desbalance (F1 macro, F1 por clase o MCC) además de la exactitud?"),
    ("QA5", "¿Compara contra una línea base o contra el estado del arte?"),
    ("QA6", "¿Reporta variabilidad (desviación, intervalos) o pruebas de significancia estadística?"),
    ("QA7", "¿Es reproducible (código, datos e hiperparámetros disponibles o descritos)?"),
    ("QA8", "¿Discute limitaciones o amenazas a la validez?"),
]

estudios = list(csv.DictReader(open(os.path.join(BASE, "incluidos.csv"), encoding="utf-8-sig")))


def encabezado(ws, columnas, grises=()):
    ws.append(columnas)
    for i, c in enumerate(ws[1], 1):
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="7F7F7F" if columnas[i - 1] in grises else AZUL)
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 45
    ws.freeze_panes = "D2"


wb = Workbook()

# --- Instrucciones ---------------------------------------------------------------
ins = wb.active
ins.title = "Instrucciones"
for t in ["Formulario de extracción y evaluación de calidad — RSL",
          "",
          "Se llena leyendo el TEXTO COMPLETO de cada estudio (PDF en Articulos\\PDF\\<carpeta>\\, empieza por el ID).",
          "",
          "Hoja 'Extraccion': una fila por estudio.",
          "  · Las columnas grises son una sugerencia automática sacada del resumen: VERIFÍCALAS con el texto completo.",
          "  · Las columnas azules se llenan a mano; las que tienen lista desplegable solo aceptan esos valores.",
          "  · Si el estudio NO cumple los criterios al leerlo completo, escribe en 'Elegible tras texto completo' = No",
          "    y el motivo (CE1–CE5). Ese número va al PRISMA.",
          "",
          "Hoja 'Calidad': responde QA1–QA8 con Sí (1), Parcial (0.5) o No (0). El total (0–8) se calcula solo.",
          "  · La calidad NO excluye estudios: se usa para ponderar la evidencia en la síntesis.",
          "",
          "Hoja 'Criterios': definición de cada pregunta de calidad y de cada campo.",
          "",
          "Reparto sugerido: el primer revisor (asistido por IA) llena todo; Andrés y Cristian revisan una",
          "muestra del 20 % de la hoja 'Calidad' y se mide el acuerdo, igual que en el cribado."]:
    ins.append([t])
ins["A1"].font = Font(bold=True, size=13)
ins.column_dimensions["A"].width = 115

# --- Extracción --------------------------------------------------------------------
ex = wb.create_sheet("Extraccion")
fijas = ["ID", "Año", "Título", "DOI", "Carpeta del PDF"]
sugeridas = ["Sugerido: enfoque", "Sugerido: técnica contra el desbalance", "Sugerido: fuente", "Sugerido: conjuntos"]
eleg = ["Elegible tras texto completo", "Motivo si no es elegible"]
cols = fijas + sugeridas + eleg + [c for c, _, _ in CAMPOS]
encabezado(ex, cols, grises=sugeridas)

def carpeta(r):
    if r["doi"].lower() in {"10.48084/etasr.14313", "10.1007/s00766-025-00451-8", "10.1109/access.2025.3570359"}:
        return "01_semillas"
    if r["aumento"]:
        return "03_aumentacion"
    return {"LLM": "05_llm", "Transformer": "04_bert_transformers"}.get(r["enfoque"], "06_ml_dl_clasico")

for r in estudios:
    ex.append([int(r["id"]), int(r["anio"]), r["titulo"], r["doi"], carpeta(r),
               r["enfoque"], r["aumento"] or ("menciona desbalance" if r["desbalance"] == "True" else ""),
               r["fuente_datos"], r["datasets"]] + [None] * (len(eleg) + len(CAMPOS)))
n = len(estudios) + 1
for col_idx, c in enumerate(cols, 1):
    letra = ex.cell(row=1, column=col_idx).column_letter
    ex.column_dimensions[letra].width = 55 if c == "Título" else 14 if c in ("ID", "Año") else 24
    if c in sugeridas:
        for fila in range(2, n + 1):
            ex.cell(row=fila, column=col_idx).fill = PatternFill("solid", fgColor=GRIS)
listas = {c: v for c, _, v in CAMPOS if v}
listas["Elegible tras texto completo"] = "Sí,No"
listas["Motivo si no es elegible"] = "CE1,CE2,CE3,CE4,CE5,Sin acceso al texto"
for col_idx, c in enumerate(cols, 1):
    if c in listas:
        dv = DataValidation(type="list", formula1=f'"{listas[c]}"', allow_blank=True)
        ex.add_data_validation(dv)
        letra = ex.cell(row=1, column=col_idx).column_letter
        dv.add(f"{letra}2:{letra}{n}")
for fila in ex.iter_rows(min_row=2):
    for c in fila:
        c.alignment = Alignment(vertical="top", wrap_text=True)
ex.auto_filter.ref = f"A1:{ex.cell(row=1, column=len(cols)).column_letter}{n}"

# --- Calidad -----------------------------------------------------------------------
ca = wb.create_sheet("Calidad")
encabezado(ca, ["ID", "Año", "Título"] + [q for q, _ in CALIDAD] + ["Total (0–8)", "Revisor", "Observaciones"])
for r in estudios:
    ca.append([int(r["id"]), int(r["anio"]), r["titulo"]] + [None] * len(CALIDAD) + [None, None, None])
dvq = DataValidation(type="list", formula1='"Sí,Parcial,No"', allow_blank=True)
ca.add_data_validation(dvq)
dvq.add(f"D2:K{n}")
for fila in range(2, n + 1):
    rango = f"D{fila}:K{fila}"
    ca.cell(row=fila, column=12).value = (f'=IF(COUNTA({rango})=0,"",COUNTIF({rango},"Sí")+0.5*COUNTIF({rango},"Parcial"))')
    for col in range(4, 12):
        ca.cell(row=fila, column=col).fill = PatternFill("solid", fgColor=AMARILLO)
ca.conditional_formatting.add(f"L2:L{n}", CellIsRule(operator="lessThan", formula=["4"], fill=PatternFill("solid", fgColor="F8CBAD")))
ca.conditional_formatting.add(f"L2:L{n}", CellIsRule(operator="greaterThanOrEqual", formula=["6"], fill=PatternFill("solid", fgColor="C6EFCE")))
for col, ancho in zip("ABC", [6, 7, 60]):
    ca.column_dimensions[col].width = ancho
for col in "DEFGHIJK":
    ca.column_dimensions[col].width = 9
ca.column_dimensions["L"].width = 11
ca.column_dimensions["N"].width = 40
for fila in ca.iter_rows(min_row=2):
    for c in fila:
        c.alignment = Alignment(vertical="top", wrap_text=c.column == 3)

# --- Criterios ----------------------------------------------------------------------
cr = wb.create_sheet("Criterios")
cr.append(["Código", "Pregunta de calidad", "Sí (1)", "Parcial (0.5)", "No (0)"])
for c in cr[1]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=AZUL)
GUIA = {
    "QA1": ("Objetivo o preguntas explícitas", "Objetivo general sin preguntas", "No se identifica"),
    "QA2": ("Origen, tamaño y distribución por clase", "Origen y tamaño, sin distribución", "No se describe"),
    "QA3": ("Protocolo descrito y aumentación solo en entrenamiento / sin duplicados", "Protocolo descrito, sin precisar la fuga", "No se describe"),
    "QA4": ("F1 macro o por clase, o MCC", "Solo F1 ponderado", "Solo exactitud"),
    "QA5": ("Compara con línea base y con el estado del arte", "Solo una de las dos", "Sin comparación"),
    "QA6": ("Prueba estadística o intervalos", "Solo desviación estándar", "Ni variabilidad ni pruebas"),
    "QA7": ("Código y datos públicos", "Hiperparámetros descritos, sin código", "No reproducible"),
    "QA8": ("Sección de amenazas o limitaciones", "Menciona alguna limitación", "No discute limitaciones"),
}
for q, t in CALIDAD:
    cr.append([q, t, *GUIA[q]])
cr.append([])
cr.append(["Campo de extracción", "Descripción", "Valores permitidos"])
for c in cr[cr.max_row]:
    c.font = Font(bold=True)
for c, d, v in CAMPOS:
    cr.append([c, d, (v or "texto libre").replace(",", ", ")])
for col, ancho in zip("ABCDE", [34, 80, 40, 36, 26]):
    cr.column_dimensions[col].width = ancho
for fila in cr.iter_rows(min_row=2):
    for c in fila:
        c.alignment = Alignment(vertical="top", wrap_text=True)

wb.save(SALIDA)
print(f"{len(estudios)} estudios -> {os.path.normpath(SALIDA)}")
