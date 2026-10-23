"""Vuelca extraccion_texto_completo.json en las hojas Extraccion y Calidad de RSL/formulario_extraccion_calidad.xlsx.

Las columnas se emparejan por el nombre de la cabecera. Solo escribe las filas de los estudios que están en el JSON;
las demás quedan como están. Para la revisión de la muestra, Andrés y Cristian anotan sus correcciones en una copia.

Uso:  python Articulos/extraccion/json_a_formulario.py
"""
import json
import os

from openpyxl import load_workbook

BASE = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.join(BASE, "..", "..", "RSL", "formulario_extraccion_calidad.xlsx")
REVISOR = "Revisor 1 (asistido por IA)"

estudios = json.load(open(os.path.join(BASE, "extraccion_texto_completo.json"), encoding="utf-8"))["estudios"]
wb = load_workbook(XLSX)

ws = wb["Extraccion"]
cab = {c.value: c.column for c in ws[1] if c.value}
for fila in ws.iter_rows(min_row=2):
    i = str(fila[0].value)
    if i not in estudios:
        continue
    datos = dict(estudios[i])
    datos.setdefault("Elegible tras texto completo", "Sí")
    for campo, valor in datos.items():
        if campo in cab and not isinstance(valor, dict):
            ws.cell(fila[0].row, cab[campo]).value = valor

wc = wb["Calidad"]
cabq = {c.value: c.column for c in wc[1] if c.value}
for fila in wc.iter_rows(min_row=2):
    i = str(fila[0].value)
    if i not in estudios or "calidad" not in estudios[i]:
        continue
    for q, v in estudios[i]["calidad"].items():
        wc.cell(fila[0].row, cabq[q]).value = {1: "Sí", 0.5: "Parcial", 0: "No"}[v]
    wc.cell(fila[0].row, cabq["Revisor"]).value = REVISOR
    wc.cell(fila[0].row, cabq["Observaciones"]).value = estudios[i].get("obs_calidad", "")

faltan = sorted(set(estudios) - {str(f[0].value) for f in ws.iter_rows(min_row=2)}, key=int)
if faltan:
    print("Sin fila en el formulario:", faltan)
wb.save(XLSX)
print(f"{len(estudios)} estudios volcados -> {os.path.normpath(XLSX)}")
