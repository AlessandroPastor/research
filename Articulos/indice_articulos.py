"""Genera Articulos/PDF/INDICE_ARTICULOS.xlsx: los 145 estudios incluidos con la ecuación de Scopus que los
recuperó, el estado de su PDF y dónde descargarlo.

Hojas:
  Estudios     un estudio por fila; columnas S1–S4 marcan la ecuación que lo encontró
  Ecuaciones   texto completo de S1–S4, para copiar en Scopus
  Por ecuación los estudios agrupados por ecuación (un estudio puede aparecer en varias)

Uso:  python Articulos/indice_articulos.py
"""
import csv
import json
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

BASE = os.path.dirname(os.path.abspath(__file__))
PDF = os.path.join(BASE, "PDF")
EJES = {"S1": "Aumentación de datos y desbalance", "S2": "BERT y transformers",
        "S3": "LLM como clasificador", "S4": "GitHub, auto-etiquetado y ampliación"}


def leer(ruta):
    with open(ruta, encoding="utf-8-sig") as fh:
        return list(csv.DictReader(fh))


formulas = [l for l in open(os.path.join(BASE, "FORMULAS_BUSQUEDA_SCOPUS.txt"), encoding="utf-8").read().splitlines()
            if l.startswith("TITLE-ABS-KEY")][:4]
busquedas = {f["id"]: f["busquedas"] for f in leer(os.path.join(BASE, "cribado.csv"))}
incluidos = leer(os.path.join(BASE, "incluidos.csv"))
faltantes = {f["id"]: f for f in leer(os.path.join(PDF, "FALTANTES.csv"))}
estudios = json.load(open(os.path.join(BASE, "extraccion", "extraccion_texto_completo.json"), encoding="utf-8"))["estudios"]

descargados = {}
for carpeta in sorted(os.listdir(PDF)):
    ruta = os.path.join(PDF, carpeta)
    if os.path.isdir(ruta):
        for nombre in os.listdir(ruta):
            ident = nombre.split("_", 1)[0]
            if nombre.lower().endswith(".pdf") and ident.isdigit():
                descargados[ident] = f"{carpeta}/{nombre}"


def estado(f):
    i = f["id"]
    if i in descargados:
        return "Descargado", descargados[i], ""
    falta = faltantes.get(i)
    doi = f["doi"]
    enlace = falta["enlace_libre"] if falta and falta["enlace_libre"] else (f"https://doi.org/{doi}" if doi else "")
    destino = f"{falta['carpeta']}/{falta['nombre_de_archivo']}" if falta else ""
    if doi.lower().startswith("10.1016/"):
        return "ScienceDirect (UPeU)", destino, enlace
    if falta and falta["acceso"].startswith("abierto"):
        return "Acceso abierto (navegador)", destino, enlace
    if not doi:
        return "Sin DOI (buscar por título)", destino, "https://scholar.google.com/scholar?q=" + f["titulo"].replace(" ", "+")
    return "De pago (sin suscripción)", destino, enlace


def sintesis(i):
    if i not in estudios:
        return ""
    return "Excluido (" + estudios[i]["Motivo si no es elegible"] + ")" if estudios[i].get(
        "Elegible tras texto completo") == "No" else "Sí"


COLOR = {"Descargado": "C6EFCE", "ScienceDirect (UPeU)": "DDEBF7", "Acceso abierto (navegador)": "E2EFDA",
         "De pago (sin suscripción)": "FCE4D6", "Sin DOI (buscar por título)": "FFF2CC"}
NEGRITA = Font(bold=True, color="FFFFFF")
CABECERA = PatternFill("solid", fgColor="305496")

wb = Workbook()
ws = wb.active
ws.title = "Estudios"
cab = ["#", "Año", "Título", "S1", "S2", "S3", "S4", "Estado del PDF", "Abrir", "Archivo (carpeta/nombre)",
       "En la síntesis", "Enfoque", "DOI"]
ws.append(cab)
filas = []
for f in sorted(incluidos, key=lambda x: int(x["id"])):
    est, archivo, enlace = estado(f)
    s = busquedas.get(f["id"], "").split()
    filas.append((f, est, archivo, enlace, s))
    ws.append([int(f["id"]), int(f["anio"]), f["titulo"]] + ["X" if k in s else "" for k in EJES]
              + [est, "abrir" if enlace else "", archivo, sintesis(f["id"]), f["enfoque"], f["doi"]])
    fila = ws.max_row
    ws.cell(fila, 8).fill = PatternFill("solid", fgColor=COLOR[est])
    if enlace:
        ws.cell(fila, 9).hyperlink = enlace
        ws.cell(fila, 9).font = Font(color="0563C1", underline="single")
for c in ws[1]:
    c.font, c.fill = NEGRITA, CABECERA
for col, ancho in zip("ABCDEFGHIJKLM", [5, 6, 70, 4, 4, 4, 4, 26, 7, 55, 13, 18, 30]):
    ws.column_dimensions[col].width = ancho
ws.freeze_panes = "D2"
ws.auto_filter.ref = ws.dimensions

we = wb.create_sheet("Ecuaciones")
we.append(["Ecuación", "Eje", "Estudios incluidos", "Descargados", "Cadena para Scopus (Advanced document search)"])
for k, (clave, eje) in enumerate(EJES.items()):
    del_eje = [x for x in filas if clave in x[4]]
    we.append([clave, eje, len(del_eje), sum(1 for x in del_eje if x[1] == "Descargado"), formulas[k]])
    we.cell(we.max_row, 5).alignment = Alignment(wrap_text=True, vertical="top")
for c in we[1]:
    c.font, c.fill = NEGRITA, CABECERA
for col, ancho in zip("ABCDE", [9, 34, 10, 12, 120]):
    we.column_dimensions[col].width = ancho

wp = wb.create_sheet("Por ecuación")
wp.append(["Ecuación", "#", "Año", "Título", "Estado del PDF", "Abrir"])
for clave, eje in EJES.items():
    for f, est, archivo, enlace, s in filas:
        if clave in s:
            wp.append([f"{clave} · {eje}", int(f["id"]), int(f["anio"]), f["titulo"], est, "abrir" if enlace else ""])
            wp.cell(wp.max_row, 5).fill = PatternFill("solid", fgColor=COLOR[est])
            if enlace:
                wp.cell(wp.max_row, 6).hyperlink = enlace
                wp.cell(wp.max_row, 6).font = Font(color="0563C1", underline="single")
for c in wp[1]:
    c.font, c.fill = NEGRITA, CABECERA
for col, ancho in zip("ABCDEF", [40, 5, 6, 80, 26, 7]):
    wp.column_dimensions[col].width = ancho
wp.freeze_panes = "A2"
wp.auto_filter.ref = wp.dimensions

salida = os.path.join(PDF, "INDICE_ARTICULOS.xlsx")
wb.save(salida)
cuenta = {}
for x in filas:
    cuenta[x[1]] = cuenta.get(x[1], 0) + 1
print(cuenta, "->", salida)
