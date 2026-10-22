"""Convierte Articulos/PDF/FALTANTES.csv en un Excel con enlaces clicables y casilla de control.

Uso:  python Articulos/faltantes_excel.py
"""
import csv
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

PDF = os.path.join(os.path.dirname(os.path.abspath(__file__)), "PDF")
filas = list(csv.DictReader(open(os.path.join(PDF, "FALTANTES.csv"), encoding="utf-8-sig")))
COLOR = {"abierto": "E2EFDA", "de pago": "FCE4D6", "sin DOI": "FFF2CC"}

wb = Workbook()
ws = wb.active
ws.title = "Faltantes"
ws.append(["#", "Acceso", "Abrir", "Guardar en la carpeta", "Con este nombre", "Título", "¿Descargado?"])
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="1F4E78")
for n, r in enumerate(filas, 1):
    enlace = r["enlace_libre"] or r["enlace"] or ("https://scholar.google.com/scholar?q=" + r["titulo"].replace(" ", "+"))
    ws.append([n, r["acceso"], "Abrir", f"Articulos\\PDF\\{r['carpeta']}", r["nombre_de_archivo"], r["titulo"], None])
    celda = ws.cell(row=n + 1, column=3)
    celda.hyperlink = enlace
    celda.font = Font(color="0563C1", underline="single")
    color = next(v for k, v in COLOR.items() if r["acceso"].startswith(k))
    ws.cell(row=n + 1, column=2).fill = PatternFill("solid", fgColor=color)
for col, ancho in zip("ABCDEFG", [5, 30, 8, 32, 52, 80, 13]):
    ws.column_dimensions[col].width = ancho
for fila in ws.iter_rows(min_row=2):
    for c in fila:
        c.alignment = Alignment(vertical="top", wrap_text=c.column == 6)
dv = DataValidation(type="list", formula1='"Sí,No se encontró"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f"G2:G{len(filas) + 1}")
ws.freeze_panes = "C2"
ws.auto_filter.ref = f"A1:G{len(filas) + 1}"

inst = wb.create_sheet("Instrucciones", 0)
for t in ["Cómo descargar los PDF que faltan",
          "",
          "1. Filtra la columna 'Acceso'. Empieza por los 'abierto': se descargan gratis, sin cuenta.",
          "2. Clic en 'Abrir'. Se abre la página del artículo; busca el botón 'PDF' o 'Download PDF'.",
          "3. Guarda el PDF en la carpeta indicada y con el nombre EXACTO de la columna 'Con este nombre'.",
          "4. Marca '¿Descargado?' = Sí. Si no lo encuentras, marca 'No se encontró'.",
          "5. Los 'de pago': entra primero con tu cuenta de la UPeU (biblioteca / acceso remoto) y luego usa 'Abrir'.",
          "6. Los 'sin DOI': el enlace abre Google Scholar con el título; descarga la versión que aparezca como [PDF].",
          "",
          "Al terminar, avisa a Claude: el script descargar_pdfs.py verifica qué archivos ya están."]:
    inst.append([t])
inst["A1"].font = Font(bold=True, size=13)
inst.column_dimensions["A"].width = 110
wb.save(os.path.join(PDF, "FALTANTES.xlsx"))
print(f"{len(filas)} filas -> Articulos/PDF/FALTANTES.xlsx")
